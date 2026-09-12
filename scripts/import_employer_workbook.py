"""Convert a reviewed employer workbook into safe discovery registries.

This imports names only. It does not probe sites, scrape providers or infer
that an employer has a public job board.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
IGNORED_SHEETS = {"All Employers", "Research Additions"}


def workbook_rows(workbook_path: Path) -> tuple[list[dict[str, object]], list[dict[str, str]]]:
    workbook = load_workbook(workbook_path, data_only=True)
    employers: list[dict[str, object]] = []
    providers: list[dict[str, str]] = []
    seen: set[str] = set()
    for sheet in workbook.worksheets:
        if sheet.title in IGNORED_SHEETS:
            continue
        values = [str(sheet.cell(row, 1).value).strip() for row in range(2, sheet.max_row + 1)
                  if isinstance(sheet.cell(row, 1).value, str) and str(sheet.cell(row, 1).value).strip()]
        if sheet.title == "Job Providers":
            providers.extend({"provider": value, "status": "research"} for value in values)
            continue
        for company in values:
            key = company.casefold()
            if key in seen:
                continue
            seen.add(key)
            employers.append({"company": company, "sector": sheet.title, "priority": 3})
    return employers, providers


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/import_employer_workbook.py path/to/employers.xlsx")
    employers, providers = workbook_rows(Path(sys.argv[1]))
    (ROOT / "data" / "sydney_employer_universe.json").write_text(
        json.dumps(employers, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (ROOT / "data" / "recruitment_source_targets.json").write_text(
        json.dumps(providers, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(f"Imported {len(employers)} employer targets and {len(providers)} recruitment-source targets.")


if __name__ == "__main__":
    main()
